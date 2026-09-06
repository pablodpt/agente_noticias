"""Diseño estratégico de una cartera UCITS no ETF para un residente en España.

El módulo trabaja exclusivamente con *categorías* de fondos. No contiene ISIN,
nombres comerciales ni instrucciones de contratación. Los pesos representan una
cartera modelo para discusión y deben contrastarse con la situación del inversor.
"""
from __future__ import annotations

import csv
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Iterable


@dataclass(frozen=True)
class PosicionModelo:
    bloque: str
    categoria: str
    peso: int
    riesgo: str
    rol: str


@dataclass(frozen=True)
class TipoFondo:
    categoria: str
    tipo: str
    perfil: str
    horizonte: str
    riesgos: str
    encaje: str


@dataclass(frozen=True)
class RevisionAnual:
    item: str
    indicador: str
    alerta: str
    accion: str


CARTERA_MODELO = (
    PosicionModelo("Core", "Global equity diversified UCITS", 14, "Medio-alto", "Motor global diversificado; limita sesgos regionales y de estilo."),
    PosicionModelo("Core", "Quality equity UCITS", 8, "Medio-alto", "Prioriza ROE, márgenes, balance y estabilidad de beneficios."),
    PosicionModelo("Core", "Low volatility equity UCITS", 7, "Medio", "Reduce beta y amortigua caídas, sin eliminar riesgo de renta variable."),
    PosicionModelo("Core", "Renta fija corporativa grado de inversión UCITS", 24, "Medio-bajo", "Renta, estabilidad y contrapeso; duración preferentemente intermedia."),
    PosicionModelo("Core", "Mixtos moderados UCITS", 15, "Medio", "Asignación dinámica y diversificación entre renta fija y variable."),
    PosicionModelo("Satélite", "Tecnología de calidad UCITS", 3, "Alto", "Crecimiento secular acotado, con énfasis en beneficios y balance."),
    PosicionModelo("Satélite", "Salud UCITS", 5, "Medio-alto", "Crecimiento defensivo y diversificación sectorial."),
    PosicionModelo("Satélite", "Infraestructuras UCITS", 5, "Medio-alto", "Flujos relativamente estables y posible sensibilidad positiva a inflación."),
    PosicionModelo("Satélite", "Renta fija flexible UCITS", 11, "Medio", "Gestiona duración, crédito y regímenes de tipos cambiantes."),
    PosicionModelo("Temático", "Inteligencia artificial UCITS con historial amplio", 5, "Alto", "Participación limitada en crecimiento estructural, evitando que domine el riesgo."),
    PosicionModelo("Temático", "Semiconductores UCITS", 3, "Alto", "Exposición acotada al habilitador tecnológico y a su ciclo."),
)

TIPOS_FONDO = (
    TipoFondo("Global equity diversified", "Indexado amplio o activo", "Crecimiento moderado", "≥ 5 años; dentro de esta cartera, revisión a 3 años", "Renta variable, divisa, concentración de índices", "Diversifica geografías y sectores y evita depender de una tesis única."),
    TipoFondo("Quality equity", "Activo sistemático o indexado por factor", "Moderado", "≥ 5 años", "Valoraciones elevadas, sesgo large-cap, factor quality", "Empresas con rentabilidad, márgenes, deuda y beneficios más robustos."),
    TipoFondo("Low volatility equity", "Indexado por factor o activo", "Defensivo", "≥ 5 años", "Quedar rezagado en rallies, sesgos sectoriales, falsa sensación de seguridad", "Busca menor beta y drawdown que la renta variable global."),
    TipoFondo("Renta fija corporativa grado de inversión", "Activo diversificado o indexado", "Defensivo", "3 años", "Duración, ampliación de spreads, impago, divisa", "Aporta carry y menor volatilidad; exige control de duración y calidad crediticia."),
    TipoFondo("Mixtos moderados", "Activo", "Moderado", "3–5 años", "Asignación errónea, costes, opacidad y solapamientos", "Delega ajustes entre activos y puede suavizar la trayectoria."),
    TipoFondo("Tecnología de calidad", "Activo", "Crecimiento", "≥ 5 años", "Valoración, concentración, regulación, sensibilidad a tipos", "Peso reducido y filtro de calidad contienen una exposición de beta alta."),
    TipoFondo("Salud", "Activo o indexado sectorial", "Moderado", "≥ 5 años", "Regulación, patentes, ensayos clínicos, concentración", "Demanda relativamente defensiva y motores estructurales de largo plazo."),
    TipoFondo("Infraestructuras", "Activo", "Moderado", "3–5 años", "Tipos, regulación, apalancamiento, sensibilidad política", "Ingresos contractuales o regulados pueden diversificar el ciclo económico."),
    TipoFondo("Renta fija flexible", "Activo", "Moderado", "3 años", "Riesgo de gestor, crédito, duración, derivados y liquidez", "Permite adaptar duración y crédito, complementando el bloque investment grade."),
    TipoFondo("Inteligencia artificial con historial amplio", "Activo", "Crecimiento", "≥ 5 años", "Tema joven, valoración, concentración y solapamiento tecnológico", "Solo como satélite pequeño, líquido y con proceso probado."),
    TipoFondo("Semiconductores", "Activo", "Crecimiento", "≥ 5 años", "Alta ciclicidad, geopolítica, capex y concentración", "Peso máximo acotado para capturar crecimiento sin dominar el drawdown."),
)

CHECKLIST = (
    RevisionAnual("Residencia y elegibilidad fiscal", "Residencia fiscal, titularidad y registro del fondo", "El fondo o el inversor no cumple el régimen de traspasos", "Solicitar validación fiscal y documental antes de cualquier cambio."),
    RevisionAnual("Pesos por bloque", "Desviación frente a Core 68 / Satélite 24 / Temático 8", "Desviación absoluta > 3 puntos por bloque", "Reequilibrar preferentemente mediante aportaciones o traspasos elegibles."),
    RevisionAnual("Riesgo total", "Volatilidad 1 y 3 años, beta y VaR", "Volatilidad > 10–12% o beta > 0,65 de forma persistente", "Revisar presupuesto de riesgo y solapamientos; no actuar por una sola observación."),
    RevisionAnual("Drawdown", "Máxima caída y tiempo de recuperación", "Drawdown > 15% o claramente peor que la referencia moderada", "Identificar contribuyentes y comprobar si la diversificación funcionó."),
    RevisionAnual("Tipos y duración", "Duración efectiva y sensibilidad a ±100 pb", "Duración > 5 años sin compensación o fuerte concentración de curva", "Revisar el rango estratégico de duración de la renta fija."),
    RevisionAnual("Crédito y liquidez", "Rating medio, high yield, spreads y liquidez", "Deterioro bajo BBB-, high yield no previsto o liquidez insuficiente", "Restaurar límites de calidad, diversificación y liquidez mediante el proceso aprobado."),
    RevisionAnual("Calidad empresarial", "ROE, margen, deuda neta/EBITDA y estabilidad de BPA", "Deterioro simultáneo durante 2 revisiones", "Reevaluar si la categoría conserva su exposición real a quality."),
    RevisionAnual("Factores", "Exposición a quality, low volatility, tamaño y value/growth", "Un factor explica > 35% del riesgo activo", "Reducir duplicidades entre categorías en el diseño estratégico."),
    RevisionAnual("Diversificación", "País, sector, emisor y top 10 subyacente", "País > 55%, sector > 25% o emisor subyacente > 5%", "Revisar solapamientos agregados, no solo cada fondo por separado."),
    RevisionAnual("Temáticos", "Peso conjunto, valoración, amplitud e historial", "Temáticos > 10%, tesis debilitada o concentración creciente", "Restablecer el límite estratégico tras revisión del comité."),
    RevisionAnual("Costes y consistencia", "TER/costes totales, tracking error, rotación y estilo", "Costes aumentan o deriva de estilo durante 2 periodos", "Comparar el vehículo con alternativas de la misma categoría, sin decisión automática."),
    RevisionAnual("Divisa", "Exposición neta y política de cobertura, especialmente en renta fija", "Divisa domina la volatilidad defensiva", "Revisar cobertura de divisa en el bloque de renta fija."),
)


def validar_modelo() -> None:
    """Comprueba restricciones estructurales declaradas por la política."""
    total = sum(p.peso for p in CARTERA_MODELO)
    bloques = {b: sum(p.peso for p in CARTERA_MODELO if p.bloque == b)
               for b in {p.bloque for p in CARTERA_MODELO}}
    assert total == 100, f"Los pesos suman {total}, no 100"
    assert 60 <= bloques.get("Core", 0) <= 70
    assert 20 <= bloques.get("Satélite", 0) <= 30
    assert 5 <= bloques.get("Temático", 0) <= 10


def contexto_macro() -> list[tuple[str, str, str]]:
    """Marco por escenarios; evita fingir una previsión puntual en tiempo real."""
    return [
        ("Inflación", "Desinflación gradual pero irregular; servicios y salarios pueden ser persistentes.", "Favorece duración intermedia, no apuestas extremas de tipos."),
        ("Tipos y bancos centrales", "Normalización dependiente de datos, con riesgo de pausas o repuntes.", "Combinar grado de inversión y renta fija flexible; controlar duración."),
        ("Ciclo económico", "Crecimiento moderado con divergencias regionales y riesgo de desaceleración.", "Sesgo a quality, salud, low volatility y amplia diversificación global."),
        ("Beneficios", "Crecimiento desigual y exigencia de convertir inversión en flujo de caja.", "Priorizar márgenes, ROE, deuda sostenible y estabilidad de beneficios."),
        ("Riesgos de cola", "Geopolítica, energía, crédito, liquidez y concentración tecnológica.", "Limitar temáticos, agregar exposiciones subyacentes y revisar drawdown."),
    ]


def _markdown_table(headers: list[str], rows: Iterable[Iterable[object]]) -> str:
    def clean(value: object) -> str:
        return str(value).replace("|", "\\|").replace("\n", " ")
    lines = ["| " + " | ".join(headers) + " |", "|" + "|".join(["---"] * len(headers)) + "|"]
    lines.extend("| " + " | ".join(clean(v) for v in row) + " |" for row in rows)
    return "\n".join(lines)


def construir_informe_markdown(fecha: date | None = None) -> str:
    """Devuelve el informe completo, apto para consola o Markdown."""
    validar_modelo()
    fecha = fecha or date.today()
    macro = _markdown_table(
        ["Variable", "Escenario central a 3 años", "Implicación de cartera"],
        contexto_macro(),
    )
    tabla1 = _markdown_table(
        ["Bloque", "Categoría de fondo UCITS", "Peso sugerido (%)", "Nivel de riesgo", "Rol en la cartera"],
        ((p.bloque, p.categoria, p.peso, p.riesgo, p.rol) for p in CARTERA_MODELO),
    )
    tabla2 = _markdown_table(
        ["Categoría", "Tipo de fondo (indexado / activo)", "Perfil", "Horizonte recomendado", "Riesgos clave", "Motivo de encaje en perfil moderado"],
        ((t.categoria, t.tipo, t.perfil, t.horizonte, t.riesgos, t.encaje) for t in TIPOS_FONDO),
    )
    tabla3 = _markdown_table(
        ["Ítem a revisar", "Indicador sugerido", "Señal de alerta", "Acción sugerida"],
        ((r.item, r.indicador, r.alerta, r.accion) for r in CHECKLIST),
    )
    return f"""# Cartera estratégica UCITS no ETF — perfil moderado

**Fecha de marco:** {fecha.isoformat()}  
**Horizonte de planificación:** 3 años, con revisión anual.  
**Asignación por bloques:** Core 68% · Satélite 24% · Temático 8%.

## Contexto macro y criterios

{macro}

El escenario no es una predicción puntual: la revisión anual debe contrastarlo con inflación, tipos reales, curvas, spreads y ciclo. El factor *quality* se evalúa mediante ROE, márgenes, apalancamiento y estabilidad de beneficios. *Low volatility* se usa para moderar beta, no como sustituto de renta fija.

## Tabla 1 — Diseño de cartera por bloques

{tabla1}

## Tabla 2 — Tipos de fondos UCITS recomendables por categoría

{tabla2}

## Tabla 3 — Checklist de revisión anual

{tabla3}

## Riesgos, oportunidades y prioridades

- **Riesgos:** horizonte corto para renta variable, inflación persistente, shocks de tipos/crédito, concentración tecnológica, divisa y solapamientos invisibles entre categorías.
- **Oportunidades:** carry de crédito de calidad, diversificación de duración, crecimiento rentable, salud e infraestructuras, y temáticas estructurales con presupuesto de riesgo limitado.
- **Prioridades anuales:** verificar primero fiscalidad y elegibilidad; después riesgo agregado, drawdown, duración/crédito, concentración subyacente, deriva factorial y costes.
- **Disciplina:** los umbrales abren una revisión; no constituyen por sí solos una orden de compra o venta.

## Fiscalidad española y alcance

Se priorizan fondos UCITS **no cotizados** potencialmente aptos para el régimen español de traspasos. La condición UCITS, por sí sola, no garantiza el diferimiento: depende del vehículo, su registro/comercialización, la residencia y el tipo de titular. Debe verificarse con la comercializadora y un profesional fiscal. Se excluyen ETF porque, con carácter general, no disfrutan del mismo régimen de traspasos en España.

Documento educativo basado en categorías, no en productos concretos. No es asesoramiento personalizado ni sugiere compras o ventas.
"""


def exportar_csvs(directorio: str | Path) -> list[Path]:
    """Exporta las tres tablas como CSV UTF-8 con BOM, cómodos para Excel."""
    validar_modelo()
    directorio = Path(directorio)
    directorio.mkdir(parents=True, exist_ok=True)
    tablas = [
        ("01_cartera_bloques.csv", ["Bloque", "Categoría de fondo UCITS", "Peso sugerido (%)", "Nivel de riesgo", "Rol en la cartera"],
         ((p.bloque, p.categoria, p.peso, p.riesgo, p.rol) for p in CARTERA_MODELO)),
        ("02_tipos_fondos.csv", ["Categoría", "Tipo de fondo (indexado / activo)", "Perfil", "Horizonte recomendado", "Riesgos clave", "Motivo de encaje en perfil moderado"],
         ((t.categoria, t.tipo, t.perfil, t.horizonte, t.riesgos, t.encaje) for t in TIPOS_FONDO)),
        ("03_checklist_anual.csv", ["Ítem a revisar", "Indicador sugerido", "Señal de alerta", "Acción sugerida"],
         ((r.item, r.indicador, r.alerta, r.accion) for r in CHECKLIST)),
    ]
    paths = []
    for nombre, headers, rows in tablas:
        path = directorio / nombre
        with path.open("w", encoding="utf-8-sig", newline="") as fh:
            writer = csv.writer(fh, delimiter=";")
            writer.writerow(headers)
            writer.writerows(rows)
        paths.append(path)
    return paths


def exportar_excel(destino: str | Path) -> Path:
    """Crea un .xlsx con contexto y las tres tablas (requiere openpyxl)."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:  # pragma: no cover - depende del entorno
        raise RuntimeError("Instala dependencias con: pip install -r requirements.txt") from exc

    validar_modelo()
    destino = Path(destino)
    destino.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    datasets = [
        ("Contexto", ["Variable", "Escenario central a 3 años", "Implicación de cartera"], contexto_macro()),
        ("Cartera", ["Bloque", "Categoría de fondo UCITS", "Peso sugerido (%)", "Nivel de riesgo", "Rol en la cartera"],
         [(p.bloque, p.categoria, p.peso, p.riesgo, p.rol) for p in CARTERA_MODELO]),
        ("Tipos de fondos", ["Categoría", "Tipo de fondo (indexado / activo)", "Perfil", "Horizonte recomendado", "Riesgos clave", "Motivo de encaje en perfil moderado"],
         [(t.categoria, t.tipo, t.perfil, t.horizonte, t.riesgos, t.encaje) for t in TIPOS_FONDO]),
        ("Revisión anual", ["Ítem a revisar", "Indicador sugerido", "Señal de alerta", "Acción sugerida"],
         [(r.item, r.indicador, r.alerta, r.accion) for r in CHECKLIST]),
    ]
    for titulo, headers, rows in datasets:
        ws = wb.create_sheet(titulo)
        ws.append(headers)
        for row in rows:
            ws.append(row)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        for idx, header in enumerate(headers, start=1):
            content_width = max([len(str(header))] + [len(str(ws.cell(r, idx).value or "")) for r in range(2, ws.max_row + 1)])
            ws.column_dimensions[get_column_letter(idx)].width = min(max(content_width + 2, 14), 48)
        if titulo == "Cartera":
            ws.column_dimensions["C"].width = 20
            ws.append(["TOTAL", "", f"=SUM(C2:C{ws.max_row})", "", ""])
            ws.cell(ws.max_row, 1).font = Font(bold=True)
            ws.cell(ws.max_row, 3).font = Font(bold=True)

    notas = wb.create_sheet("Notas")
    notas["A1"] = "Alcance y fiscalidad"
    notas["A1"].font = Font(bold=True, color="FFFFFF")
    notas["A1"].fill = PatternFill("solid", fgColor="1F4E78")
    notas["A2"] = (f"Fecha del marco: {date.today().isoformat()}. Cartera modelo educativa por categorías. "
                   "No recomienda productos, compras ni ventas. Se priorizan fondos UCITS no cotizados, "
                   "pero UCITS no garantiza por sí solo el régimen español de traspasos: verificar registro, "
                   "comercialización, residencia y titularidad. 'Indexado' significa fondo indexado no "
                   "cotizado, nunca ETF. Horizonte 3 años y revisión anual.")
    notas["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    notas.column_dimensions["A"].width = 120
    notas.row_dimensions[2].height = 90
    wb.save(destino)
    return destino


def main() -> None:
    """CLI autónoma; Markdown y CSV solo necesitan la biblioteca estándar."""
    import argparse

    parser = argparse.ArgumentParser(description="Cartera modelo UCITS no ETF por categorías")
    parser.add_argument("formato", nargs="?", default="markdown",
                        choices=["markdown", "csv", "excel"])
    parser.add_argument("--salida", default=None,
                        help="Archivo .xlsx o directorio CSV de destino")
    args = parser.parse_args()
    if args.formato == "markdown":
        print(construir_informe_markdown())
    elif args.formato == "csv":
        rutas = exportar_csvs(args.salida or "tablas_ucits")
        print("\n".join(f"Tabla exportada: {ruta}" for ruta in rutas))
    else:
        ruta = exportar_excel(args.salida or "informe_ucits.xlsx")
        print(f"Informe UCITS exportado: {ruta}")


if __name__ == "__main__":
    main()
